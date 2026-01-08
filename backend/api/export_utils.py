"""
Utilitários para exportação de dados em PDF e Excel
"""
from io import BytesIO
from django.http import HttpResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime


def export_to_excel(queryset, filename, columns, sheet_name='Dados'):
    """
    Exporta queryset para Excel
    
    Args:
        queryset: QuerySet do Django
        filename: Nome do arquivo
        columns: Lista de tuplas (campo, título)
        sheet_name: Nome da planilha
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    # Estilo do cabeçalho
    header_fill = PatternFill(start_color='F4A100', end_color='F4A100', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Adiciona cabeçalhos
    for col_num, (field, title) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Adiciona dados
    for row_num, obj in enumerate(queryset, 2):
        for col_num, (field, title) in enumerate(columns, 1):
            cell = ws.cell(row=row_num, column=col_num)
            
            # Obtém valor do campo (suporta campos aninhados)
            value = obj
            for part in field.split('.'):
                if hasattr(value, part):
                    value = getattr(value, part)
                elif isinstance(value, dict):
                    value = value.get(part)
                else:
                    value = ''
                    break
            
            # Formata valor
            if callable(value):
                value = value()
            
            cell.value = str(value) if value is not None else ''
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')
    
    # Ajusta largura das colunas
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Cria resposta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response


def export_to_pdf(data, filename, title, columns, orientation='portrait'):
    """
    Exporta dados para PDF
    
    Args:
        data: Lista de dicionários com os dados
        filename: Nome do arquivo
        title: Título do relatório
        columns: Lista de tuplas (campo, título)
        orientation: 'portrait' ou 'landscape'
    """
    buffer = BytesIO()
    
    # Define tamanho da página
    pagesize = A4 if orientation == 'portrait' else (A4[1], A4[0])
    doc = SimpleDocTemplate(buffer, pagesize=pagesize, topMargin=0.5*inch, bottomMargin=0.5*inch)
    
    # Estilos
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#F4A100'),
        alignment=TA_CENTER,
        spaceAfter=30,
    )
    
    # Elementos do documento
    elements = []
    
    # Título
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 0.2*inch))
    
    # Data de geração
    date_style = ParagraphStyle(
        'DateStyle',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.grey,
        alignment=TA_CENTER,
    )
    elements.append(Paragraph(f'Gerado em: {datetime.now().strftime("%d/%m/%Y %H:%M")}', date_style))
    elements.append(Spacer(1, 0.3*inch))
    
    # Prepara dados da tabela
    table_data = [[title for field, title in columns]]
    
    for item in data:
        row = []
        for field, title in columns:
            value = item.get(field, '')
            row.append(str(value) if value is not None else '')
        table_data.append(row)
    
    # Cria tabela
    table = Table(table_data)
    table.setStyle(TableStyle([
        # Cabeçalho
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F4A100')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        
        # Corpo
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FFF8E1')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    
    elements.append(table)
    
    # Rodapé
    elements.append(Spacer(1, 0.5*inch))
    footer_style = ParagraphStyle(
        'Footer',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.grey,
        alignment=TA_CENTER,
    )
    elements.append(Paragraph('Smart-Hive - Sistema de Gerenciamento de Apiário', footer_style))
    
    # Gera PDF
    doc.build(elements)
    
    # Cria resposta HTTP
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response
